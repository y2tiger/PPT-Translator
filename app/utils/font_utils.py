"""Font utilities for PPT processing."""
import subprocess
import structlog
from pathlib import Path
from pptx import Presentation
from pptx.shapes.group import GroupShape
from pptx.util import Pt
from pptx.dml.color import RGBColor

logger = structlog.get_logger(__name__)

# Default font for Korean text (Pretendard is modern and clean)
DEFAULT_KOREAN_FONT = "Pretendard"

# Available fonts for user selection
AVAILABLE_FONTS = {
    "pretendard": {
        "name": "Pretendard",
        "display_name": "Pretendard (프리텐다드)",
        "description": "현대적이고 깔끔한 한글 폰트",
    },
    "nanum_gothic": {
        "name": "Nanum Gothic",
        "display_name": "Nanum Gothic (나눔고딕)",
        "description": "널리 사용되는 무료 한글 폰트",
    },
    "nanum_myeongjo": {
        "name": "Nanum Myeongjo",
        "display_name": "Nanum Myeongjo (나눔명조)",
        "description": "명조체 스타일 한글 폰트",
    },
    "noto_sans_kr": {
        "name": "Noto Sans CJK KR",
        "display_name": "Noto Sans KR (노토 산스)",
        "description": "Google의 다국어 지원 폰트",
    },
    "trebuchet_ms": {
        "name": "Trebuchet MS",
        "display_name": "Trebuchet MS",
        "description": "깔끔한 산세리프 영문 폰트",
    },
    "original": {
        "name": None,  # Keep original font
        "display_name": "원본 유지",
        "description": "PPT의 원본 폰트 유지 (폰트가 없으면 대체됨)",
    },
}

# Mapping of common Korean fonts to Pretendard (default)
FONT_MAPPING = {
    # Windows Korean fonts → Pretendard (default)
    "맑은 고딕": DEFAULT_KOREAN_FONT,
    "Malgun Gothic": DEFAULT_KOREAN_FONT,
    "굴림": DEFAULT_KOREAN_FONT,
    "Gulim": DEFAULT_KOREAN_FONT,
    "돋움": DEFAULT_KOREAN_FONT,
    "Dotum": DEFAULT_KOREAN_FONT,
    "바탕": "Nanum Myeongjo",  # Keep serif for serif
    "Batang": "Nanum Myeongjo",
    "궁서": "Nanum Myeongjo",
    "Gungsuh": "Nanum Myeongjo",
    "새굴림": DEFAULT_KOREAN_FONT,
    "New Gulim": DEFAULT_KOREAN_FONT,
    # Common presentation fonts
    "Arial": DEFAULT_KOREAN_FONT,
    "Calibri": DEFAULT_KOREAN_FONT,
    "Times New Roman": "Nanum Myeongjo",
    "Verdana": DEFAULT_KOREAN_FONT,
    # Keep these as-is (usually available)
    "Pretendard": "Pretendard",
    "Nanum Gothic": "Nanum Gothic",
    "Nanum Myeongjo": "Nanum Myeongjo",
    "Noto Sans KR": "Noto Sans CJK KR",
    "Noto Sans CJK KR": "Noto Sans CJK KR",
}


def get_available_fonts() -> list[dict]:
    """Get list of available fonts for UI selection."""
    return [
        {
            "id": font_id,
            "name": info["name"],
            "display_name": info["display_name"],
            "description": info["description"],
        }
        for font_id, info in AVAILABLE_FONTS.items()
    ]


def get_font_mapping(target_font: str | None = None) -> dict:
    """
    Get font mapping with optional custom target font.

    Args:
        target_font: Target font name (e.g., "Pretendard", "Nanum Gothic")
                    If None, use default mapping.
    """
    if target_font is None:
        return FONT_MAPPING.copy()

    # Create custom mapping with user-selected font
    custom_mapping = {}
    for original, default_replacement in FONT_MAPPING.items():
        # For sans-serif fonts, use target_font
        # For serif fonts (바탕, 궁서, Times), keep serif alternative
        if default_replacement in ["Nanum Myeongjo"]:
            custom_mapping[original] = default_replacement
        else:
            custom_mapping[original] = target_font

    return custom_mapping


def extract_fonts_from_ppt(ppt_path: str) -> set[str]:
    """Extract all font names used in a PPT file."""
    fonts = set()

    try:
        prs = Presentation(ppt_path)

        for slide in prs.slides:
            for shape in slide.shapes:
                fonts.update(_extract_fonts_from_shape(shape))

        logger.info("fonts_extracted", ppt_path=ppt_path, fonts=list(fonts))
        return fonts

    except Exception as e:
        logger.error("font_extraction_failed", error=str(e))
        return set()


def _extract_fonts_from_shape(shape, depth: int = 0) -> set[str]:
    """Extract fonts from a single shape (recursive for groups)."""
    fonts = set()

    # Handle grouped shapes (with depth limit to prevent infinite loops)
    if isinstance(shape, GroupShape) and depth < 10:
        try:
            for child_shape in shape.shapes:
                fonts.update(_extract_fonts_from_shape(child_shape, depth + 1))
        except Exception:
            pass
        return fonts

    # Handle text frames
    if hasattr(shape, "text_frame"):
        try:
            for paragraph in shape.text_frame.paragraphs:
                # Paragraph-level font
                if paragraph.font and paragraph.font.name:
                    fonts.add(paragraph.font.name)

                # Run-level fonts
                for run in paragraph.runs:
                    if run.font and run.font.name:
                        fonts.add(run.font.name)
        except Exception:
            pass

    # Handle tables
    if hasattr(shape, "table"):
        try:
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        if paragraph.font and paragraph.font.name:
                            fonts.add(paragraph.font.name)
                        for run in paragraph.runs:
                            if run.font and run.font.name:
                                fonts.add(run.font.name)
        except Exception:
            pass

    return fonts


def apply_font_to_ppt(ppt_path: str, output_path: str, target_font: str) -> bool:
    """
    Apply a specific font to all text in a PPT file.

    Args:
        ppt_path: Path to input PPT
        output_path: Path to save modified PPT
        target_font: Font name to apply

    Returns:
        True if successful, False otherwise
    """
    try:
        prs = Presentation(ppt_path)

        for slide in prs.slides:
            for shape in slide.shapes:
                _apply_font_to_shape(shape, target_font)

        prs.save(output_path)
        logger.info("font_applied_to_ppt", target_font=target_font, output_path=output_path)
        return True

    except Exception as e:
        logger.error("font_application_failed", error=str(e))
        return False


def _apply_font_to_shape(shape, target_font: str, depth: int = 0):
    """Apply font to a single shape (recursive for groups)."""
    # Handle grouped shapes (with depth limit)
    if isinstance(shape, GroupShape) and depth < 10:
        try:
            for child_shape in shape.shapes:
                _apply_font_to_shape(child_shape, target_font, depth + 1)
        except Exception:
            pass
        return

    # Handle text frames
    if hasattr(shape, "text_frame"):
        try:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font:
                        run.font.name = target_font
        except Exception:
            pass

    # Handle tables
    if hasattr(shape, "table"):
        try:
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font:
                                run.font.name = target_font
        except Exception:
            pass


def get_system_fonts() -> set[str]:
    """Get list of fonts available on the system."""
    try:
        result = subprocess.run(
            ["fc-list", "--format=%{family}\n"],
            capture_output=True,
            text=True,
            timeout=30
        )

        if result.returncode == 0:
            fonts = set()
            for line in result.stdout.strip().split("\n"):
                # fc-list may return comma-separated family names
                for font in line.split(","):
                    fonts.add(font.strip())
            return fonts

    except Exception as e:
        logger.warning("system_font_list_failed", error=str(e))

    return set()


def check_missing_fonts(ppt_path: str, target_font: str | None = None) -> dict:
    """
    Check which fonts from the PPT are missing on the system.

    Returns:
        dict with:
            - used_fonts: fonts used in PPT
            - available_fonts: fonts available on system
            - missing_fonts: fonts that are missing
            - suggested_mappings: suggested font replacements
            - target_font: font that will be used for translation
    """
    used_fonts = extract_fonts_from_ppt(ppt_path)
    system_fonts = get_system_fonts()
    font_mapping = get_font_mapping(target_font)

    missing_fonts = set()
    suggested_mappings = {}

    for font in used_fonts:
        if font and font not in system_fonts:
            # Check if there's a known mapping
            mapped_font = font_mapping.get(font)
            if mapped_font and mapped_font in system_fonts:
                suggested_mappings[font] = mapped_font
            else:
                missing_fonts.add(font)
                # Default to Pretendard for unknown fonts
                suggested_mappings[font] = DEFAULT_KOREAN_FONT

    result = {
        "used_fonts": list(used_fonts),
        "available_fonts": len(system_fonts),
        "missing_fonts": list(missing_fonts),
        "suggested_mappings": suggested_mappings,
        "target_font": target_font or DEFAULT_KOREAN_FONT,
    }

    logger.info("font_check_complete", **result)
    return result


def create_libreoffice_font_substitution(target_font: str | None = None) -> str:
    """
    Create LibreOffice font substitution configuration.
    Returns the path to the configuration file.
    """
    # LibreOffice user profile path
    lo_profile = Path.home() / ".config" / "libreoffice" / "4" / "user"
    lo_profile.mkdir(parents=True, exist_ok=True)

    registrymodifications_path = lo_profile / "registrymodifications.xcu"

    # Get font mapping with optional custom target
    font_mapping = get_font_mapping(target_font)

    # Create font substitution entries
    substitutions = []
    for original, replacement in font_mapping.items():
        substitutions.append(f'<item oor:path="/org.openoffice.Office.Common/Font/Substitution/FontPairs">'
                           f'<node oor:name="{original}" oor:op="replace">'
                           f'<prop oor:name="ReplaceFont"><value>{replacement}</value></prop>'
                           f'<prop oor:name="Always"><value>true</value></prop>'
                           f'</node></item>')

    config_content = f'''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
{''.join(substitutions)}
</oor:items>
'''

    try:
        registrymodifications_path.write_text(config_content)
        logger.info("libreoffice_font_config_created", path=str(registrymodifications_path))
        return str(registrymodifications_path)
    except Exception as e:
        logger.error("libreoffice_font_config_failed", error=str(e))
        return ""
